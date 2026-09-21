"""Выгрузка таблиц ручного поиска ТТХ в текст: источник для разбора в data/ttx_manual.csv.

Excel читается только здесь, дальше по конвейеру идёт текст, поэтому разбор
не зависит от openpyxl и воспроизводится без бинарных файлов.
"""

from __future__ import annotations

import pathlib

import openpyxl

SRC = pathlib.Path(__file__).resolve().parents[2] / "resources" / "Ручной поиск"
OUT = pathlib.Path(__file__).resolve().parents[2] / "data" / "manual_dump"


def cell(v: object) -> str:
    if v is None:
        return ""
    return str(v).replace("\r", " ").replace("\n", " ⏎ ").strip()


def dump(path: pathlib.Path) -> pathlib.Path:
    wb = openpyxl.load_workbook(path, data_only=True)
    lines = [f"ФАЙЛ: {path.name}", ""]
    for ws in wb.worksheets:
        lines.append(f"### ЛИСТ «{ws.title}» — {ws.max_row} строк × {ws.max_column} колонок")
        lines.append("")
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [cell(c) for c in row]
            while cells and not cells[-1]:
                cells.pop()
            if not cells:
                continue
            lines.append(f"--- строка {i} ---")
            # Пары «колонка = значение» вместо таблицы: 82 колонки в ряд нечитаемы
            for j, c in enumerate(cells, start=1):
                if c:
                    lines.append(f"  [{j}] {c}")
            lines.append("")
    target = OUT / f"{path.stem}.txt"
    target.write_text("\n".join(lines), encoding="utf-8")
    return target


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for path in sorted(SRC.glob("*.xlsx")):
        target = dump(path)
        print(f"{path.name} → {target.relative_to(target.parents[2])} ({target.stat().st_size} байт)")


if __name__ == "__main__":
    main()
